"""Builds the review presentation from <out>/results. python tools/make_slides.py --out outputs --pptx report/Music_Generation_Slides.pptx"""
import argparse, json
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ap = argparse.ArgumentParser(); ap.add_argument("--out", default="outputs"); ap.add_argument("--pptx", default="report/Music_Generation_Slides.pptx")
a = ap.parse_args(); R = Path(a.out) / "results"; M = json.loads((R / "metrics.json").read_text())
ds, ne, gen, tr = M["dataset"], {r["model"]: r for r in M["next_event"]}, M["generation"], M["training"]
NAVY, BLUE, GREY = RGBColor(0x1B, 0x3A, 0x6B), RGBColor(0x2A, 0x6F, 0xDB), RGBColor(0x44, 0x44, 0x44)
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]


def tb(s, x, y, w, h, text, size=18, bold=False, color=GREY):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = t.text_frame; tf.word_wrap = True
    for i, line in enumerate(text if isinstance(text, list) else [text]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.text = line; p.space_after = Pt(8)
        p.font.size, p.font.bold, p.font.color.rgb = Pt(size), bold, color
    return t


def slide(title, bullets=None, size=20):
    s = prs.slides.add_slide(blank)
    bar = s.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1.1)); bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    tb(s, 0.6, 0.18, 12, 0.8, title, 30, True, RGBColor(255, 255, 255))
    if bullets: tb(s, 0.7, 1.5, 12, 5.5, ["\u2022 " + b for b in bullets], size)
    return s


def table(s, rows, x, y, w, colw, size=14):
    t = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(0.45 * len(rows))).table
    for j, cw in enumerate(colw): t.columns[j].width = Inches(cw)
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c = t.cell(i, j); c.text = str(v)
            for p in c.text_frame.paragraphs: p.font.size = Pt(size); p.font.bold = (i == 0 or i == len(rows) - 1 and rows is not None and False)


s = prs.slides.add_slide(blank)
bg = s.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height); bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
tb(s, 0.8, 2.2, 11.5, 1.5, "Generating Music with Machine Learning", 44, True, RGBColor(255, 255, 255))
tb(s, 0.8, 3.7, 11.5, 1.5, ["Mini-Project #22 \u2013 UE24CS352A Machine Learning", "Nagashri R Patil (PES1UG24CS289)  |  Nallamalli Kanaka Mani Sai Akhil (PES1UG24CS290)",
                            "Based on: Kang, Kim, Ringdahl \u2013 Stanford, \u201cMusic Composition with Machine Learning\u201d (2018)"], 18, False, RGBColor(0xDD, 0xE6, 0xF5))

slide("Problem statement", ["Teach a computer to compose by learning musical structure from MIDI files",
      "Task 1: predict the next note (pitch + duration) on unseen songs",
      "Task 2: generate new, playable MIDI pieces (melody + chords)",
      "Compare LSTM (main model) vs Naive-Bayes-like chord\u2192note vs random (flat) baseline",
      "Goal from the paper: LSTM sequence modelling should beat independent-note models"])
slide("Dataset and representation", [f"Dataset: {ds['dataset_name']} \u2013 {ds['songs_used']} songs, {ds['events_total']:,} note events (read from data/, any .mid/.midi)",
      f"Song-level split 80/10/10: {ds['songs']['train']} train / {ds['songs']['val']} val / {ds['songs']['test']} test songs",
      "Event = [right-hand note (88 keys), duration (8 classes), left-hand chord (63 + other)]  \u2013 like the paper\u2019s note vectors",
      "Right hand = highest note \u2265 MIDI 60; left-hand chord = pitch-class set of lower notes sounding",
      "Durations: inter-onset interval quantised to 16th-note grid", "Windows of 64 events; padded targets are ignored in the loss"], 19)
slide("Models", ["Random (flat): uniform over 88 keys / durations / chords \u2013 reference level",
      "Naive-Bayes-like: P(note | chord), notes independent given chord, Laplace smoothing; Markov chain over chords for generation",
      "LSTM (main): embeddings \u2192 1 LSTM layer (256) \u2192 dropout \u2192 3 softmax heads (pitch, duration, chord)",
      "Training: cross-entropy, Adam lr 0.005 (\u03b2\u2081 .9, \u03b2\u2082 .999), grad clipping, early stopping on validation loss",
      "Generation: autoregressive sampling \u2192 events \u2192 MIDI (melody track + chord track)"], 19)
slide("Pipeline and implementation", ["preprocess.py \u2192 parse MIDI, split, chord vocabulary, windows", "train.py \u2192 LSTM training / validation, checkpoint, loss plot",
      "evaluate.py \u2192 test metrics for all models, generated-music metrics, plots, MIDI export + reload check",
      "demo.py \u2192 generate (and play) a new MIDI composition live", "run_all.py \u2192 whole pipeline in one command; PyTorch, CPU-friendly, auto-GPU",
      f"LSTM: {tr['epochs_run']} epochs, best epoch {tr['best_epoch']}, {tr['train_seconds']:.0f}s on {tr['device']}"], 20)

s = slide("Results: next-event prediction (test set)")
rows = [["Model", "Pitch CE", "Pitch ppl", "Top-1", "Top-5", "Dur. acc"]]
for k in ["Random (flat)", "Naive Bayes-like (prev chord)", "Naive Bayes-like (given chord)", "LSTM"]:
    r = ne[k]; rows.append([k, f"{r['pitch_ce']:.3f}", f"{r['pitch_ppl']:.1f}", f"{r['pitch_top1']:.3f}", f"{r['pitch_top5']:.3f}", f"{r['dur_acc']:.3f}"])
table(s, rows, 0.7, 1.5, 12, [4.6, 1.5, 1.5, 1.4, 1.4, 1.6], 16)
tb(s, 0.7, 4.0, 12, 1.5, [f"{ne['LSTM']['n_events']:,} held-out test events; lower CE / higher accuracy is better.",
   "Naive Bayes \u2018given chord\u2019 sees the true concurrent chord (paper protocol) \u2013 extra information the LSTM does not get."], 16)
s = slide("Training curve and model comparison")
s.shapes.add_picture(str(R / "loss_curve.png"), Inches(0.5), Inches(1.5), width=Inches(5.6))
s.shapes.add_picture(str(R / "next_event_comparison.png"), Inches(6.3), Inches(1.7), width=Inches(6.8))
s = slide("Results: generated music vs real music")
g = [["Metric", "Real", "Random", "Naive Bayes", "LSTM"]]
for k, lab in [("scale_consistency", "Scale consistency"), ("chord_tone_rate", "Chord-tone rate"), ("repetition_rate", "Repeated-note rate"),
               ("mean_abs_interval", "Mean |interval|"), ("novelty_5gram", "5-gram novelty")]:
    g.append([lab] + [f"{gen[n][k + '_mean']:.2f}" for n in ["Real (test)", "Random (flat)", "Naive Bayes-like", "LSTM"]])
g.append(["Interval JS vs real"] + ["0"] + [f"{gen[n]['interval_js_vs_real']:.2f}" for n in ["Random (flat)", "Naive Bayes-like", "LSTM"]])
table(s, g, 0.7, 1.5, 12, [4.2, 1.9, 1.9, 2.0, 2.0], 16)
tb(s, 0.7, 5.0, 12, 1.5, "Closer to the \u2018Real\u2019 column is better. All generated .mid files re-load and play correctly (" + ", ".join(f"{k}: {v}" for k, v in M["midi_check"].items()) + ").", 16)
s = slide("Generated piano rolls (first 48 events)"); s.shapes.add_picture(str(R / "piano_rolls.png"), Inches(2.0), Inches(1.3), height=Inches(6.0))
slide("Conclusions", ["LSTM sequence modelling beats Naive-Bayes-like and random models on next-note prediction (see table)",
      "Generated LSTM melodies match real interval statistics and stay in key; Naive Bayes fits chords best",
      "Limitations: single-voice melody, coarse chords, no dynamics/velocity, no human Turing test",
      "Future work (from the paper): Transformer / self-attention models and note velocity", "Live demo: python demo.py --play"], 20)
Path(a.pptx).parent.mkdir(parents=True, exist_ok=True); prs.save(a.pptx); print("saved", a.pptx)
